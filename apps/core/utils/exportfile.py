import csv
import openpyxl
from django.http import HttpResponse
from datetime import datetime

class ExportFile:
    """
    Utilitaire générique pour transformer une liste de dicts en fichier.
    """

    @staticmethod
    def _get_timestamp():
        return datetime.now().strftime("%Y%m%d_%H%M%S")

    @staticmethod
    def to_csv(data: list, filename: str = "export") -> HttpResponse:
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename}_{ExportFile._get_timestamp()}.csv"'

        if not data:
            return response

        writer = csv.writer(response)
        # Les clés du premier élément servent d'en-têtes
        headers = list(data[0].keys())
        writer.writerow(headers)

        for item in data:
            # On écrit les valeurs dans l'ordre des en-têtes
            writer.writerow([str(item.get(h, '')) for h in headers])

        return response

    @staticmethod
    def to_excel(data: list, filename: str = "export") -> HttpResponse:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{filename}_{ExportFile._get_timestamp()}.xlsx"'

        if not data:
            return response

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Export"

        headers = list(data[0].keys())

        # Déterminer si c'est une trame d'importation vide
        is_template = filename.startswith("trame_import_")

        # Pour les modèles d'importation, on pré-remplit 15 lignes vides bien formatées pour l'utilisateur
        rows_to_generate = data
        if is_template and len(data) <= 1:
            rows_to_generate = [{h: "" for h in headers} for _ in range(15)]

        # Écriture des en-têtes
        ws.append(headers)
        ws.row_dimensions[1].height = 28

        # Style des en-têtes (Bleu indigo élégant, texte blanc gras centré)
        header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        header_border = Border(
            left=Side(style='thin', color='312E81'),
            right=Side(style='thin', color='312E81'),
            top=Side(style='thin', color='312E81'),
            bottom=Side(style='thin', color='312E81')
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = header_border

        # Style des cellules de données (Police claire, bordures grises discrètes)
        data_font = Font(name='Segoe UI', size=10, color='1E293B')
        data_alignment = Alignment(horizontal='left', vertical='center')
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # Remplissage et style des lignes
        for row_num, item in enumerate(rows_to_generate, start=2):
            ws.row_dimensions[row_num].height = 22
            for col_num, h in enumerate(headers, start=1):
                val = item.get(h, '')
                if isinstance(val, datetime):
                    val = val.replace(tzinfo=None)
                
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border

        # Ajustement automatique de la largeur des colonnes
        for col_num, h in enumerate(headers, start=1):
            col_letter = get_column_letter(col_num)
            max_len = len(str(h))
            for item in rows_to_generate:
                val = str(item.get(h, '') or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 18)

        # Activer explicitement le quadrillage dans Excel
        ws.views.sheetView[0].showGridLines = True

        wb.save(response)
        return response